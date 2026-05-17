from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # 1. Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "DataSoul"
    subtitle.text = "The AI-Augmented Data Intelligence Platform"
    
    # 2. The Problem
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "The Problem"
    tf = body_shape.text_frame
    tf.text = "Messy Data & Manual Effort"
    p = tf.add_paragraph()
    p.text = "Modern data workflows are plagued by messy datasets, inconsistent formatting, and manual cleaning bottlenecks."
    p.level = 1
    p2 = tf.add_paragraph()
    p2.text = "Data privacy concerns limit the use of cloud LLMs for sensitive business data."
    p2.level = 1

    # 3. The DataSoul Solution
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "The DataSoul Solution"
    tf = body_shape.text_frame
    tf.text = "An automated, privacy-first data pipeline."
    p = tf.add_paragraph()
    p.text = "Integrates local LLMs (Ollama) to intelligently profile, clean, and analyze tabular data with zero external data leakage."
    p.level = 1

    # 4. Core Capabilities
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Core Capabilities"
    tf = body_shape.text_frame
    tf.text = "Automated Data Cleaning: Intelligent imputation, deduplication, and anomaly resolution."
    p = tf.add_paragraph()
    p.text = "RAG-Augmented Insights: Chat with your data using Retrieval-Augmented Generation for natural, expert-level executive narratives."
    p = tf.add_paragraph()
    p.text = "Continuous Profiling: Real-time threat detection and data quality scoring (aiming for 100% quality scores)."

    # 5. Architecture Overview
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Architecture Overview"
    tf = body_shape.text_frame
    tf.text = "Frontend: Interactive dashboard for monitoring data health and interacting with the data."
    p = tf.add_paragraph()
    p.text = "Backend: FastAPI-driven engine with specialized micro-components (pipeline_engine, threat_detector, csv_corrector, narrative_engine)."
    p = tf.add_paragraph()
    p.text = "Intelligence Layer: Ollama (Local LLM) combined with a local RAG vector store for secure, low-latency processing."

    # 6. The Pipeline Workflow
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "The Pipeline Workflow"
    tf = body_shape.text_frame
    tf.text = "1. Ingestion & Profiling: Upload CSVs; immediate statistical profiling."
    p = tf.add_paragraph()
    p.text = "2. Threat Detection: Identify anomalies, missing values, and inconsistencies."
    p = tf.add_paragraph()
    p.text = "3. Iterative Correction: LLM-guided automatic resolution of identified threats."
    p = tf.add_paragraph()
    p.text = "4. Chat & Narrative: Generate human-readable reports and answer specific queries securely."

    # 7. Technical Differentiators
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Technical Differentiators"
    tf = body_shape.text_frame
    tf.text = "Local, secure processing ensures privacy."
    p = tf.add_paragraph()
    p.text = "Robust iteration handling for complex datasets."
    p = tf.add_paragraph()
    p.text = "Extensible architecture designed for scale."

    # 8. Conclusion
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Thank You"
    tf = body_shape.text_frame
    tf.text = "Summary: DataSoul brings AI intelligence to your messy data securely and privately."
    p = tf.add_paragraph()
    p.text = "Q&A"

    prs.save('DataSoul_Presentation.pptx')
    print("Successfully generated DataSoul_Presentation.pptx")

if __name__ == '__main__':
    create_presentation()
